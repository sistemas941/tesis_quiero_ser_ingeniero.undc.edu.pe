# Plan de tesis de prueba en Word, con titulos, listas, tablas y referencias (para test/v15_ui.js)
from docx import Document
from pathlib import Path
d = Document()
for t in ['UNIVERSIDAD NACIONAL DE CAÑETE', 'FACULTAD DE INGENIERÍA', 'PROYECTO DE TESIS',
          'REDES NEURONALES CONVOLUCIONALES PARA CLASIFICAR NEUMONÍA EN RADIOGRAFÍAS PEDIÁTRICAS, CAÑETE 2026',
          'AUTOR: Pérez Soto, Ana', 'ASESOR: Dr. Juan Ramos']:
    d.add_paragraph(t)
d.add_heading('I. PLANTEAMIENTO DEL PROBLEMA', 1)
d.add_heading('1.1 Descripción de la realidad problemática', 2)
d.add_paragraph('En los hospitales de Cañete la lectura de radiografías pediátricas depende de un solo radiólogo, lo que retrasa el diagnóstico de neumonía. En 2025 el tiempo medio de informe superó las 48 horas y varios casos se derivaron tarde.')
d.add_paragraph('Escribí este planteamiento a partir de entrevistas con el personal del hospital y de los registros del servicio de imágenes.')
d.add_heading('1.2 Formulación del problema', 2)
d.add_heading('1.2.1 Problema general', 3)
d.add_paragraph('¿En qué medida las redes neuronales convolucionales mejoran la clasificación de tipos de neumonía en radiografías de pacientes pediátricos de Cañete en 2026?')
d.add_heading('1.2.2 Problemas específicos', 3)
for t in ['¿Cuál es la exactitud del modelo frente al radiólogo?', '¿Cuál es la sensibilidad del modelo para neumonía viral?']:
    d.add_paragraph(t, style='List Number')
d.add_heading('1.3 Objetivos', 2)
d.add_heading('1.3.1 Objetivo general', 3)
d.add_paragraph('Determinar en qué medida las redes neuronales convolucionales mejoran la clasificación de tipos de neumonía en radiografías pediátricas.')
d.add_heading('1.4 Justificación', 2)
d.add_heading('1.4.1 Justificación práctica', 3)
d.add_paragraph('El estudio reduce el tiempo de diagnóstico en el hospital y beneficia a las familias de la provincia.')
d.add_heading('II. MARCO TEÓRICO', 1)
d.add_heading('2.1 Antecedentes', 2)
d.add_paragraph('Ozturk et al. (2020) propusieron DarkCovidNet para radiografías de tórax.')
d.add_heading('2.2 Definición de términos básicos', 2)
d.add_paragraph('Radiografía: imagen obtenida con rayos X del tórax del paciente.')
d.add_paragraph('Neumonía: infección que inflama los alvéolos pulmonares.')
d.add_heading('III. HIPÓTESIS Y VARIABLES', 1)
d.add_heading('3.1 Operacionalización de las variables', 2)
d.add_paragraph('Tabla 1')
d.add_paragraph('Operacionalización de las variables')
t = d.add_table(rows=3, cols=4)
for i, fila in enumerate([['Variable', 'Dimensión', 'Indicador', 'Escala'], ['Redes neuronales', 'Arquitectura', 'Número de capas', 'Razón'], ['', 'Entrenamiento', 'Épocas', 'Razón']]):
    for j, v in enumerate(fila): t.cell(i, j).text = v
d.add_heading('IV. METODOLOGÍA', 1)
d.add_heading('4.1 Tipo de investigación', 2)
d.add_paragraph('La investigación tiene enfoque cuantitativo, es de tipo aplicada y de nivel explicativo.')
d.add_heading('4.2 Población y muestra', 2)
d.add_paragraph('La población son 1 200 radiografías pediátricas del hospital entre 2024 y 2025.')
d.add_heading('V. CRONOGRAMA', 1)
t = d.add_table(rows=3, cols=5)
for i, fila in enumerate([['Actividad', 'M1', 'M2', 'M3', 'M4'], ['Revisión de literatura', 'X', 'X', '', ''], ['Entrenamiento del modelo', '', '', 'X', 'X']]):
    for j, v in enumerate(fila): t.cell(i, j).text = v
d.add_heading('VI. PRESUPUESTO', 1)
t = d.add_table(rows=4, cols=4)
for i, fila in enumerate([['Descripción', 'Cantidad', 'Costo unitario', 'Total'], ['Asesor estadístico', '1', '800', '800'], ['Servicio de internet', '6', '60', '360'], ['Total', '', '', '1160']]):
    for j, v in enumerate(fila): t.cell(i, j).text = v
d.add_heading('VII. REFERENCIAS', 1)
for r in ['Ozturk, T., Talo, M., Yildirim, E. A., Baloglu, U. B., Yildirim, O., & Acharya, U. R. (2020). Automated detection of COVID-19 cases using deep neural networks with X-ray images. Computers in Biology and Medicine, 121, 103792. https://doi.org/10.1016/j.compbiomed.2020.103792',
          'Kermany, D. S., Goldbaum, M., & Cai, W. (2018). Identifying medical diagnoses and treatable diseases by image-based deep learning. Cell, 172(5), 1122-1131.',
          'Quispe Rojas, L. (2022). Clasificación de radiografías con aprendizaje profundo [Tesis de pregrado, Universidad Nacional de Cañete]. Repositorio institucional.']:
    d.add_paragraph(r)
d.add_heading('ANEXOS', 1)
d.add_paragraph('Matriz de consistencia')
d.save(Path(__file__).with_name('plan_prueba.docx'))
print('ok')
